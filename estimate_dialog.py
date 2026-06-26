import sys
import math
import re
from datetime import datetime
from PyQt6.QtWidgets import *
from PyQt6.QtCore import *
from PyQt6.QtGui import *
from PyQt6.QtPrintSupport import QPrinter, QPrintDialog
import db

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False

try:
    from weasyprint import HTML
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False


class EstimateDialog(QDialog):
    def __init__(self, materials, client_name="Не выбран", client_address="Не указан",
                 parent=None, cutting_data=None, client_id=None, current_user=None, paint_enabled=False):
        super().__init__(parent)
        self.materials = materials
        self.client_name = client_name
        self.client_address = client_address
        self.cutting_data = cutting_data
        self.client_id = client_id
        self.current_user = current_user
        self.paint_enabled = paint_enabled
        self.setWindowTitle("Смета")
        self.setMinimumSize(1100, 600)
        self.init_ui()
        self.load_data()

    def init_ui(self):
        layout = QVBoxLayout(self)

        info_layout = QHBoxLayout()
        info_layout.addWidget(QLabel(f"<b>Клиент:</b> {self.client_name}"))
        info_layout.addStretch()
        info_layout.addWidget(QLabel(f"<b>Адрес:</b> {self.client_address}"))
        info_layout.addStretch()
        info_layout.addWidget(QLabel(f"<b>Дата:</b> {datetime.now().strftime('%d.%m.%Y')}"))
        layout.addLayout(info_layout)

        hint = QLabel(
            "💡 Совет: можно менять длину и цену прямо в таблице. Выделите строку и нажмите 'Удалить' для удаления позиции.\n"
            "📦 Комплекты добавляются как одна позиция, вложенные позиции показаны для информации."
        )
        hint.setStyleSheet("color: #666; font-style: italic; padding: 3px;")
        hint.setWordWrap(True)
        layout.addWidget(hint)

        self.table = QTableWidget()
        self.table.setColumnCount(8)
        self.table.setHorizontalHeaderLabels(
            ["№", "Наименование", "Длина/Кол-во", "Хлысты", "Ед.", "Цена (руб)", "Сумма (руб)", "Вес (кг)"]
        )
        self.table.horizontalHeader().setStretchLastSection(True)
        self.table.setColumnWidth(0, 30)
        self.table.setColumnWidth(1, 250)
        self.table.setColumnWidth(2, 100)
        self.table.setColumnWidth(3, 180)
        self.table.setColumnWidth(4, 50)
        self.table.setColumnWidth(5, 90)
        self.table.setColumnWidth(6, 90)
        self.table.setColumnWidth(7, 70)
        layout.addWidget(self.table)

        total_layout = QHBoxLayout()
        self.total_label = QLabel("<b>ИТОГО: 0.00 руб</b>")
        self.total_label.setStyleSheet("font-size: 14px; color: #0055AA;")
        total_layout.addWidget(self.total_label)

        self.total_weight_label = QLabel("<b>Вес: 0.00 кг</b>")
        self.total_weight_label.setStyleSheet("font-size: 12px; color: #555;")
        total_layout.addWidget(self.total_weight_label)

        total_layout.addStretch()
        layout.addLayout(total_layout)

        btn_layout = QHBoxLayout()

        self.add_btn = QPushButton("➕ Добавить из справочника")
        self.add_btn.clicked.connect(self.add_item_from_catalog)
        btn_layout.addWidget(self.add_btn)

        self.delete_btn = QPushButton("🗑️ Удалить позицию")
        self.delete_btn.setStyleSheet("background-color: #f44336; color: white;")
        self.delete_btn.clicked.connect(self.delete_selected_item)
        btn_layout.addWidget(self.delete_btn)

        self.refresh_btn = QPushButton("🔄 Обновить из расчёта")
        self.refresh_btn.clicked.connect(self.refresh_from_materials)
        btn_layout.addWidget(self.refresh_btn)

        self.save_btn = QPushButton("💾 Сохранить в БД")
        self.save_btn.clicked.connect(self.save_estimate)
        btn_layout.addWidget(self.save_btn)

        self.export_excel_btn = QPushButton("📊 Excel")
        self.export_excel_btn.setToolTip("Экспорт сметы в Excel")
        self.export_excel_btn.clicked.connect(self.export_to_excel)
        self.export_excel_btn.setStyleSheet("background-color: #2E7D32; color: white;")
        btn_layout.addWidget(self.export_excel_btn)

        self.export_pdf_btn = QPushButton("📄 PDF")
        self.export_pdf_btn.setToolTip("Экспорт сметы в PDF (WEasyPrint)")
        self.export_pdf_btn.clicked.connect(self.export_to_pdf)
        self.export_pdf_btn.setStyleSheet("background-color: #C62828; color: white;")
        btn_layout.addWidget(self.export_pdf_btn)

        self.print_btn = QPushButton("🖨️ Печать")
        self.print_btn.clicked.connect(self.print_pdf)
        btn_layout.addWidget(self.print_btn)

        self.close_btn = QPushButton("❌ Закрыть")
        self.close_btn.clicked.connect(self.accept)
        btn_layout.addWidget(self.close_btn)

        layout.addLayout(btn_layout)

    def delete_selected_item(self):
        row = self.table.currentRow()
        if row < 0:
            QMessageBox.warning(self, "Ошибка", "Выберите строку для удаления!")
            return

        name_item = self.table.item(row, 1)
        if not name_item:
            return

        name = name_item.text()
        is_kit_row = name.startswith("📦")
        is_child_row = name.startswith("  ├──")
        kit_group_id = name_item.data(Qt.ItemDataRole.UserRole + 1)

        if is_kit_row:
            kit_rows = 1
            row_num = row + 1
            while row_num < self.table.rowCount():
                next_name_item = self.table.item(row_num, 1)
                if not next_name_item:
                    break
                next_name = next_name_item.text()
                if next_name.startswith("  ├──"):
                    kit_rows += 1
                    row_num += 1
                else:
                    break

            reply = QMessageBox.question(
                self, "Подтверждение",
                f"Удалить комплект '{name}' и все его позиции ({kit_rows - 1} позиций)?",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
            )
            if reply == QMessageBox.StandardButton.Yes:
                for _ in range(kit_rows):
                    self.table.removeRow(row)
                self.update_row_numbers()
                self.recalc_totals()

        elif is_child_row and kit_group_id:
            reply = QMessageBox.question(
                self, "Подтверждение",
                f"Удалить позицию '{name}' из комплекта?",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
            )
            if reply == QMessageBox.StandardButton.Yes:
                self.table.removeRow(row)
                has_children = False
                for r in range(self.table.rowCount()):
                    r_name_item = self.table.item(r, 1)
                    if not r_name_item:
                        continue
                    r_kit_group_id = r_name_item.data(Qt.ItemDataRole.UserRole + 1)
                    if r_kit_group_id == kit_group_id:
                        r_num_item = self.table.item(r, 0)
                        if not r_num_item or not r_num_item.text():
                            has_children = True
                            break

                if not has_children:
                    for r in range(self.table.rowCount()):
                        r_name_item = self.table.item(r, 1)
                        if not r_name_item:
                            continue
                        r_kit_group_id = r_name_item.data(Qt.ItemDataRole.UserRole + 1)
                        if r_kit_group_id == kit_group_id:
                            self.table.removeRow(r)
                            break

                self.update_row_numbers()
                self.recalc_totals()

        else:
            reply = QMessageBox.question(
                self, "Подтверждение",
                f"Удалить позицию '{name}'?",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
            )
            if reply == QMessageBox.StandardButton.Yes:
                self.table.removeRow(row)
                self.update_row_numbers()
                self.recalc_totals()

    def load_data(self):
        if not self.materials:
            self.table.setRowCount(0)
            self.update_totals(0, 0)
            return

        groups = {}
        for mat in self.materials:
            if isinstance(mat, dict):
                profile = mat.get('profile', 'Другое')
                length_mm = float(mat.get('length_mm', 0))
                length_for_payment_mm = float(mat.get('length_for_payment_mm', length_mm))
                sticks = int(mat.get('sticks', 1))
                remainder_mm = float(mat.get('remainder_mm', 0))
                remainder_for_client = float(mat.get('remainder_for_client', 0))
                weight = float(mat.get('weight', 0))
                price = mat.get('price')
                unit = mat.get('unit', 'м.п.')
                is_paint = mat.get('is_paint', False)
                is_kit = mat.get('is_kit', False)
                kit_items = mat.get('kit_items', [])
                name = mat.get('name', profile)

                if price is not None:
                    try:
                        price = float(price)
                    except (ValueError, TypeError):
                        price = None
            else:
                try:
                    length_mm = float(mat[2]) if len(mat) > 2 else 0
                except (ValueError, TypeError):
                    length_mm = 0
                length_for_payment_mm = length_mm
                try:
                    sticks = int(float(mat[3])) if len(mat) > 3 else 1
                except (ValueError, TypeError):
                    sticks = 1
                remainder_mm = (sticks * 6000) - length_mm
                if remainder_mm < 0:
                    remainder_mm = 0
                remainder_for_client = 0
                try:
                    weight = float(mat[4]) if len(mat) > 4 else 0
                except (ValueError, TypeError):
                    weight = 0
                price = None
                if len(mat) > 5 and mat[5] is not None and mat[5] != "":
                    try:
                        price = float(mat[5])
                    except (ValueError, TypeError):
                        price = None
                profile = mat[1] if len(mat) > 1 and mat[1] else (mat[0] if len(mat) > 0 else "Другое")
                unit = 'м.п.'
                is_paint = False
                is_kit = False
                kit_items = []
                name = mat[0] if len(mat) > 0 else profile

            if profile not in groups:
                groups[profile] = {
                    'length_mm': 0,
                    'length_for_payment_mm': 0,
                    'weight': 0,
                    'sticks': 0,
                    'remainder_mm': 0,
                    'remainder_for_client': 0,
                    'price': price,
                    'unit': unit,
                    'is_paint': is_paint,
                    'is_kit': is_kit,
                    'kit_items': kit_items,
                    'name': name
                }

            groups[profile]['length_mm'] += length_mm
            groups[profile]['length_for_payment_mm'] += length_for_payment_mm
            groups[profile]['weight'] += weight
            groups[profile]['sticks'] += sticks
            groups[profile]['remainder_mm'] += remainder_mm
            groups[profile]['remainder_for_client'] += remainder_for_client
            if groups[profile]['price'] is None and price is not None:
                groups[profile]['price'] = price

        self._fill_table_from_groups(groups)

    def _fill_table_from_groups(self, groups):
        self.table.blockSignals(True)
        try:
            total_rows = len(groups)
            for data in groups.values():
                if data.get('is_kit', False) and data.get('kit_items'):
                    total_rows += len(data['kit_items'])

            self.table.setRowCount(total_rows)
            total_sum = 0.0
            total_weight = 0.0
            row = 0

            for profile, data in groups.items():
                self.table.setItem(row, 0, QTableWidgetItem(str(row + 1)))
                self.table.item(row, 0).setTextAlignment(Qt.AlignmentFlag.AlignCenter)

                name = data.get('name', profile)
                if data.get('is_kit', False):
                    name = f"📦 {name}"

                name_item = QTableWidgetItem(name)
                name_item.setFlags(Qt.ItemFlag.ItemIsEnabled | Qt.ItemFlag.ItemIsSelectable)
                if data.get('is_kit', False):
                    name_item.setBackground(QColor(200, 255, 200))
                self.table.setItem(row, 1, name_item)

                length_for_payment_m = data['length_for_payment_mm'] / 1000.0
                length_item = QTableWidgetItem(f"{length_for_payment_m:.2f}")
                length_item.setFlags(Qt.ItemFlag.ItemIsEnabled | Qt.ItemFlag.ItemIsSelectable)
                length_item.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
                self.table.setItem(row, 2, length_item)

                sticks = data['sticks']
                remainder_mm = data['remainder_mm']
                remainder_for_client = data.get('remainder_for_client', 0)
                is_paint = data.get('is_paint', False)
                is_kit = data.get('is_kit', False)

                if is_kit:
                    qty_text = "комплект"
                elif is_paint:
                    qty_text = "—"
                elif remainder_mm > 0:
                    qty_text = f"{sticks} хлыст(ов), остаток {remainder_mm:.0f} мм"
                    if remainder_for_client > 0:
                        qty_text += f" (клиенту {remainder_for_client:.0f} мм)"
                else:
                    qty_text = f"{sticks} хлыст(ов)"

                qty_item = QTableWidgetItem(qty_text)
                qty_item.setData(Qt.ItemDataRole.UserRole, sticks)
                qty_item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
                self.table.setItem(row, 3, qty_item)

                unit = data.get('unit', 'м.п.')
                unit_item = QTableWidgetItem(unit)
                unit_item.setFlags(Qt.ItemFlag.ItemIsEnabled | Qt.ItemFlag.ItemIsSelectable)
                unit_item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
                self.table.setItem(row, 4, unit_item)

                price = data['price'] if data['price'] is not None and data['price'] > 0 else 0.0
                price_item = QTableWidgetItem(f"{price:.2f}")
                price_item.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
                self.table.setItem(row, 5, price_item)

                if is_kit:
                    total = price
                elif is_paint:
                    paint_area = data['length_for_payment_mm'] / 1000.0
                    total = paint_area * price
                else:
                    total = length_for_payment_m * price

                total_item = QTableWidgetItem(f"{total:.2f}")
                total_item.setFlags(Qt.ItemFlag.ItemIsEnabled | Qt.ItemFlag.ItemIsSelectable)
                total_item.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
                self.table.setItem(row, 6, total_item)

                weight_item = QTableWidgetItem(f"{data['weight']:.2f}")
                weight_item.setFlags(Qt.ItemFlag.ItemIsEnabled | Qt.ItemFlag.ItemIsSelectable)
                weight_item.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
                self.table.setItem(row, 7, weight_item)

                if data.get('is_kit', False):
                    import time
                    kit_group_id = f"kit_{int(time.time() * 1000)}_{profile}"
                    name_item.setData(Qt.ItemDataRole.UserRole + 1, kit_group_id)

                total_sum += total
                total_weight += data['weight']
                row += 1

                if data.get('is_kit', False) and data.get('kit_items'):
                    kit_group_id = name_item.data(Qt.ItemDataRole.UserRole + 1)
                    for kit_item in data['kit_items']:
                        self.table.setItem(row, 0, QTableWidgetItem(""))

                        child_name_item = QTableWidgetItem(f"  ├── {kit_item['material_name']}")
                        child_name_item.setFlags(Qt.ItemFlag.ItemIsEnabled | Qt.ItemFlag.ItemIsSelectable)
                        child_name_item.setBackground(QColor(240, 240, 240))
                        self.table.setItem(row, 1, child_name_item)

                        qty_item = QTableWidgetItem(f"{kit_item['quantity']:.2f}")
                        qty_item.setFlags(Qt.ItemFlag.ItemIsEnabled | Qt.ItemFlag.ItemIsSelectable)
                        qty_item.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
                        self.table.setItem(row, 2, qty_item)

                        self.table.setItem(row, 3, QTableWidgetItem(""))

                        unit_item = QTableWidgetItem(kit_item.get('unit', kit_item.get('material_unit', 'шт')))
                        unit_item.setFlags(Qt.ItemFlag.ItemIsEnabled | Qt.ItemFlag.ItemIsSelectable)
                        unit_item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
                        self.table.setItem(row, 4, unit_item)

                        price_item = QTableWidgetItem(f"{kit_item['material_price']:.2f}")
                        price_item.setFlags(Qt.ItemFlag.ItemIsEnabled | Qt.ItemFlag.ItemIsSelectable)
                        price_item.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
                        self.table.setItem(row, 5, price_item)

                        subtotal = kit_item['quantity'] * kit_item['material_price']
                        subtotal_item = QTableWidgetItem(f"{subtotal:.2f}")
                        subtotal_item.setFlags(Qt.ItemFlag.ItemIsEnabled | Qt.ItemFlag.ItemIsSelectable)
                        subtotal_item.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
                        subtotal_item.setForeground(QColor(100, 100, 100))
                        self.table.setItem(row, 6, subtotal_item)

                        self.table.setItem(row, 7, QTableWidgetItem(""))

                        child_name_item.setData(Qt.ItemDataRole.UserRole + 1, kit_group_id)

                        row += 1

        finally:
            self.table.blockSignals(False)

        try:
            self.table.cellChanged.disconnect(self.on_cell_changed)
        except TypeError:
            pass
        self.table.cellChanged.connect(self.on_cell_changed)

        self.update_totals(total_sum, total_weight)

    def on_cell_changed(self, row, col):
        if col not in [2, 5]:
            return

        try:
            self.table.blockSignals(True)
            try:
                unit = self.table.item(row, 4).text() if self.table.item(row, 4) else "м.п."
                price = float(self.table.item(row, 5).text()) if self.table.item(row, 5) else 0.0

                if unit == 'м.п.':
                    try:
                        new_length_m = float(self.table.item(row, 2).text())
                    except ValueError:
                        return
                    new_length_mm = new_length_m * 1000
                    sticks = math.ceil(new_length_mm / 6000)
                    remainder = (sticks * 6000) - new_length_mm
                    remainder_for_client = remainder % 1000
                    if remainder > 0:
                        qty_text = f"{sticks} хлыст(ов), остаток {remainder:.0f} мм"
                        if remainder_for_client > 0:
                            qty_text += f" (клиенту {remainder_for_client:.0f} мм)"
                    else:
                        qty_text = f"{sticks} хлыст(ов)"
                    self.table.item(row, 3).setText(qty_text)
                    self.table.item(row, 3).setData(Qt.ItemDataRole.UserRole, sticks)
                    total = new_length_m * price
                elif unit == 'м²':
                    try:
                        new_area = float(self.table.item(row, 2).text())
                    except ValueError:
                        return
                    self.table.item(row, 3).setText("1")
                    self.table.item(row, 3).setData(Qt.ItemDataRole.UserRole, 1)
                    total = new_area * price
                else:
                    try:
                        new_qty = float(self.table.item(row, 2).text())
                    except ValueError:
                        return
                    qty_text = f"{new_qty:.0f} {unit}"
                    self.table.item(row, 3).setText(qty_text)
                    self.table.item(row, 3).setData(Qt.ItemDataRole.UserRole, new_qty)
                    total = new_qty * price

                self.table.item(row, 6).setText(f"{total:.2f}")
            finally:
                self.table.blockSignals(False)
            self.recalc_totals()
        except (ValueError, AttributeError) as e:
            print(f"Ошибка при пересчёте: {e}")

    def recalc_totals(self):
        total_sum = 0.0
        total_weight = 0.0

        kit_groups = {}

        for row in range(self.table.rowCount()):
            try:
                num_item = self.table.item(row, 0)
                name_item = self.table.item(row, 1)
                if not name_item:
                    continue

                kit_group_id = name_item.data(Qt.ItemDataRole.UserRole + 1)
                is_kit_row = name_item.text().startswith("📦")
                is_child_row = name_item.text().startswith("  ├──")

                if is_child_row and kit_group_id:
                    qty_item = self.table.item(row, 2)
                    price_item = self.table.item(row, 5)
                    if qty_item and price_item:
                        try:
                            subtotal = float(qty_item.text()) * float(price_item.text())
                            if kit_group_id not in kit_groups:
                                kit_groups[kit_group_id] = {'sum': 0, 'row': None}
                            kit_groups[kit_group_id]['sum'] += subtotal
                        except (ValueError, TypeError):
                            pass

                elif is_kit_row and kit_group_id:
                    if kit_group_id not in kit_groups:
                        kit_groups[kit_group_id] = {'sum': 0, 'row': row}
                    else:
                        kit_groups[kit_group_id]['row'] = row

                elif not kit_group_id and num_item and num_item.text():
                    total_item = self.table.item(row, 6)
                    if total_item:
                        total_sum += float(total_item.text())

                    weight_item = self.table.item(row, 7)
                    if weight_item and weight_item.text():
                        total_weight += float(weight_item.text())

            except (ValueError, AttributeError) as e:
                print(f"Ошибка при пересчёте строки {row}: {e}")
                continue

        for group_id, data in kit_groups.items():
            if data['row'] is not None:
                total_item = self.table.item(data['row'], 6)
                if total_item:
                    total_item.setText(f"{data['sum']:.2f}")
                total_sum += data['sum']

        self.update_totals(total_sum, total_weight)

    def update_totals(self, total_sum, total_weight):
        self.total_label.setText(f"<b>ИТОГО: {total_sum:.2f} руб</b>")
        self.total_weight_label.setText(f"<b>Вес: {total_weight:.2f} кг</b>")

    def add_item_from_catalog(self):
        dialog = SelectMaterialTreeDialog(self)
        if dialog.exec() == QDialog.DialogCode.Accepted:
            selected = dialog.get_selected_data()
            if selected:
                if selected.get('type') == 'kit':
                    self.add_kit_to_table(selected)
                else:
                    self.add_material_to_table(selected)

    def add_kit_to_table(self, kit_data):
        kit = db.get_kit_by_id(kit_data['id'])
        if not kit:
            QMessageBox.warning(self, "Ошибка", "Комплект не найден!")
            return

        import time
        kit_group_id = f"kit_{int(time.time() * 1000)}_{kit['id']}"

        kit_items = []
        total_price = 0
        for item in kit['items']:
            subtotal = item['quantity'] * item['material_price']
            kit_items.append({
                'material_name': item['material_name'],
                'quantity': item['quantity'],
                'unit': item.get('unit', item.get('material_unit', 'шт')),
                'material_price': item['material_price'],
                'subtotal': subtotal
            })
            total_price += subtotal

        row = self.table.rowCount()
        self.table.insertRow(row)

        self.table.blockSignals(True)
        try:
            self.table.setItem(row, 0, QTableWidgetItem(str(row + 1)))
            self.table.item(row, 0).setTextAlignment(Qt.AlignmentFlag.AlignCenter)

            name_item = QTableWidgetItem(f"📦 {kit['name']}")
            name_item.setFlags(Qt.ItemFlag.ItemIsEnabled | Qt.ItemFlag.ItemIsSelectable)
            name_item.setBackground(QColor(200, 255, 200))
            self.table.setItem(row, 1, name_item)

            self.table.setItem(row, 2, QTableWidgetItem("1"))
            self.table.setItem(row, 3, QTableWidgetItem("комплект"))

            unit_item = QTableWidgetItem("шт")
            unit_item.setFlags(Qt.ItemFlag.ItemIsEnabled | Qt.ItemFlag.ItemIsSelectable)
            unit_item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
            self.table.setItem(row, 4, unit_item)

            price_item = QTableWidgetItem(f"{total_price:.2f}")
            price_item.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
            self.table.setItem(row, 5, price_item)

            total_item = QTableWidgetItem(f"{total_price:.2f}")
            total_item.setFlags(Qt.ItemFlag.ItemIsEnabled | Qt.ItemFlag.ItemIsSelectable)
            total_item.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
            self.table.setItem(row, 6, total_item)

            weight_item = QTableWidgetItem("0.00")
            weight_item.setFlags(Qt.ItemFlag.ItemIsEnabled | Qt.ItemFlag.ItemIsSelectable)
            weight_item.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
            self.table.setItem(row, 7, weight_item)

            name_item.setData(Qt.ItemDataRole.UserRole + 1, kit_group_id)

            for kit_item in kit_items:
                row += 1
                self.table.insertRow(row)

                self.table.setItem(row, 0, QTableWidgetItem(""))

                name_item = QTableWidgetItem(f"  ├── {kit_item['material_name']}")
                name_item.setFlags(Qt.ItemFlag.ItemIsEnabled | Qt.ItemFlag.ItemIsSelectable)
                name_item.setBackground(QColor(240, 240, 240))
                self.table.setItem(row, 1, name_item)

                qty_item = QTableWidgetItem(f"{kit_item['quantity']:.2f}")
                qty_item.setFlags(Qt.ItemFlag.ItemIsEnabled | Qt.ItemFlag.ItemIsSelectable)
                qty_item.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
                self.table.setItem(row, 2, qty_item)

                self.table.setItem(row, 3, QTableWidgetItem(""))

                unit_item = QTableWidgetItem(kit_item['unit'])
                unit_item.setFlags(Qt.ItemFlag.ItemIsEnabled | Qt.ItemFlag.ItemIsSelectable)
                unit_item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
                self.table.setItem(row, 4, unit_item)

                price_item = QTableWidgetItem(f"{kit_item['material_price']:.2f}")
                price_item.setFlags(Qt.ItemFlag.ItemIsEnabled | Qt.ItemFlag.ItemIsSelectable)
                price_item.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
                self.table.setItem(row, 5, price_item)

                subtotal_item = QTableWidgetItem(f"{kit_item['subtotal']:.2f}")
                subtotal_item.setFlags(Qt.ItemFlag.ItemIsEnabled | Qt.ItemFlag.ItemIsSelectable)
                subtotal_item.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
                subtotal_item.setForeground(QColor(100, 100, 100))
                self.table.setItem(row, 6, subtotal_item)

                self.table.setItem(row, 7, QTableWidgetItem(""))

                name_item.setData(Qt.ItemDataRole.UserRole + 1, kit_group_id)

        finally:
            self.table.blockSignals(False)

        self.update_row_numbers()
        self.recalc_totals()

    def add_material_to_table(self, material, value=1.0):
        row = self.table.rowCount()
        self.table.insertRow(row)

        self.table.blockSignals(True)
        try:
            self.table.setItem(row, 0, QTableWidgetItem(str(row + 1)))
            self.table.item(row, 0).setTextAlignment(Qt.AlignmentFlag.AlignCenter)

            name_item = QTableWidgetItem(material['name'])
            name_item.setFlags(Qt.ItemFlag.ItemIsEnabled | Qt.ItemFlag.ItemIsSelectable)
            self.table.setItem(row, 1, name_item)

            unit = material.get('unit', 'м.п.')

            if unit == 'м.п.':
                length_item = QTableWidgetItem(f"{value:.2f}")
                length_item.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
                self.table.setItem(row, 2, length_item)

                length_mm = value * 1000
                sticks = math.ceil(length_mm / 6000)
                remainder = (sticks * 6000) - length_mm
                remainder_for_client = remainder % 1000
                if remainder > 0:
                    qty_text = f"{sticks} хлыст(ов), остаток {remainder:.0f} мм"
                    if remainder_for_client > 0:
                        qty_text += f" (клиенту {remainder_for_client:.0f} мм)"
                else:
                    qty_text = f"{sticks} хлыст(ов)"
                qty_item = QTableWidgetItem(qty_text)
                qty_item.setData(Qt.ItemDataRole.UserRole, sticks)
                total = value * float(material.get('retail_price', 0) or 0)
            elif unit == 'м²':
                length_item = QTableWidgetItem(f"{value:.2f}")
                length_item.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
                self.table.setItem(row, 2, length_item)
                qty_item = QTableWidgetItem("1")
                qty_item.setData(Qt.ItemDataRole.UserRole, 1)
                total = value * float(material.get('retail_price', 0) or 0)
            else:
                length_item = QTableWidgetItem(f"{value:.0f}")
                length_item.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
                self.table.setItem(row, 2, length_item)
                qty_text = f"{value:.0f} {unit}"
                qty_item = QTableWidgetItem(qty_text)
                qty_item.setData(Qt.ItemDataRole.UserRole, value)
                total = value * float(material.get('retail_price', 0) or 0)

            qty_item.setFlags(Qt.ItemFlag.ItemIsEnabled | Qt.ItemFlag.ItemIsSelectable)
            qty_item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
            self.table.setItem(row, 3, qty_item)

            unit_item = QTableWidgetItem(unit)
            unit_item.setFlags(Qt.ItemFlag.ItemIsEnabled | Qt.ItemFlag.ItemIsSelectable)
            unit_item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
            self.table.setItem(row, 4, unit_item)

            price = float(material.get('retail_price', 0) or 0)
            price_item = QTableWidgetItem(f"{price:.2f}")
            price_item.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
            self.table.setItem(row, 5, price_item)

            total_item = QTableWidgetItem(f"{total:.2f}")
            total_item.setFlags(Qt.ItemFlag.ItemIsEnabled | Qt.ItemFlag.ItemIsSelectable)
            total_item.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
            self.table.setItem(row, 6, total_item)

            weight_item = QTableWidgetItem("0.00")
            weight_item.setFlags(Qt.ItemFlag.ItemIsEnabled | Qt.ItemFlag.ItemIsSelectable)
            weight_item.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
            self.table.setItem(row, 7, weight_item)
        finally:
            self.table.blockSignals(False)

        self.update_row_numbers()
        self.recalc_totals()

    def update_row_numbers(self):
        row_num = 1
        for row in range(self.table.rowCount()):
            item = self.table.item(row, 0)
            if item and item.text():
                item.setText(str(row_num))
                row_num += 1

    def refresh_from_materials(self):
        reply = QMessageBox.question(
            self, "Подтверждение",
            "Обновить смету из текущего расчёта? Все ручные изменения будут потеряны.",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
        )
        if reply == QMessageBox.StandardButton.Yes:
            self.load_data()

    def save_estimate(self):
        items = []

        all_rows = []
        for row in range(self.table.rowCount()):
            try:
                name_item = self.table.item(row, 1)
                if not name_item:
                    continue

                num_item = self.table.item(row, 0)
                name = name_item.text()
                is_child_row = name.startswith("  ├──")

                if not is_child_row and (not num_item or not num_item.text()):
                    continue

                is_kit_row = name.startswith("📦")
                kit_group_id = name_item.data(Qt.ItemDataRole.UserRole + 1)

                unit = self.table.item(row, 4).text() if self.table.item(row, 4) else "м.п."

                try:
                    price = float(self.table.item(row, 5).text()) if self.table.item(row, 5) else 0
                except (ValueError, AttributeError):
                    price = 0.0

                try:
                    total = float(self.table.item(row, 6).text()) if self.table.item(row, 6) else 0
                except (ValueError, AttributeError):
                    total = 0.0

                try:
                    weight = float(self.table.item(row, 7).text()) if self.table.item(row, 7) else 0
                except (ValueError, AttributeError):
                    weight = 0.0

                try:
                    length_val = float(self.table.item(row, 2).text()) if self.table.item(row, 2) else 0
                except (ValueError, AttributeError):
                    length_val = 0.0

                is_paint = 'Покраска' in name

                if is_paint:
                    length_mm = length_val * 1000
                elif unit == 'м.п.':
                    length_mm = length_val * 1000
                else:
                    length_mm = 0

                qty_item = self.table.item(row, 3)
                quantity = 1
                if qty_item:
                    qty_data = qty_item.data(Qt.ItemDataRole.UserRole)
                    if qty_data is not None:
                        try:
                            quantity = int(qty_data)
                        except (ValueError, TypeError):
                            quantity = 1
                    else:
                        qty_text = qty_item.text()
                        if qty_text and qty_text != '—':
                            import re
                            match = re.search(r'(\d+)', qty_text)
                            if match:
                                quantity = int(match.group(1))

                all_rows.append({
                    'name': name,
                    'clean_name': name.replace("📦 ", "").replace("  ├── ", "").strip(),
                    'is_kit_row': is_kit_row,
                    'is_child_row': is_child_row,
                    'is_paint': is_paint,
                    'kit_group_id': kit_group_id,
                    'unit': unit,
                    'price': price,
                    'total': total,
                    'weight': weight,
                    'length_mm': length_mm,
                    'length_val': length_val,
                    'quantity': quantity
                })

            except Exception as e:
                print(f"Ошибка при сборе строки {row}: {e}")
                continue

        kit_indices = {}

        for rd in all_rows:
            if rd['is_child_row']:
                continue

            if rd['is_kit_row']:
                items.append({
                    'name': rd['clean_name'],
                    'length': 0,
                    'quantity': 1,
                    'unit': 'шт',
                    'price': rd['total'],
                    'total': rd['total'],
                    'weight': 0,
                    'note': '',
                    'is_kit': True,
                    'kit_parent_id': None,
                    'is_paint': False
                })
                kit_indices[rd['kit_group_id']] = len(items) - 1
            else:
                items.append({
                    'name': rd['clean_name'],
                    'length': rd['length_mm'],
                    'quantity': rd['quantity'],
                    'unit': rd['unit'],
                    'price': rd['price'],
                    'total': rd['total'],
                    'weight': rd['weight'],
                    'note': '',
                    'is_kit': False,
                    'kit_parent_id': None,
                    'is_paint': rd['is_paint']
                })

        for rd in all_rows:
            if not rd['is_child_row']:
                continue

            parent_idx = kit_indices.get(rd['kit_group_id'])

            items.append({
                'name': rd['clean_name'],
                'length': 0,
                'quantity': rd['length_val'],
                'unit': rd['unit'],
                'price': rd['price'],
                'total': rd['total'],
                'weight': 0,
                'note': '',
                'is_kit': False,
                'kit_parent_id': parent_idx,
                'is_paint': False
            })

        total_sum = sum(it['total'] for it in items if it.get('kit_parent_id') is None)

        estimate_data = {
            'number': f"СМ-{datetime.now().strftime('%Y%m%d')}-001",
            'date': datetime.now().strftime('%Y-%m-%d'),
            'total_amount': total_sum
        }

        estimate_id = db.save_estimate_with_kit(
            None, estimate_data, items,
            client_id=self.client_id,
            client_name=self.client_name,
            client_address=self.client_address
        )
        if estimate_id:
            QMessageBox.information(self, "Готово", f"Смета сохранена (ID: {estimate_id})")
        else:
            QMessageBox.warning(self, "Ошибка", "Не удалось сохранить смету!")

    def _collect_table_data(self):
        rows = []
        total_sum = 0.0
        total_weight = 0.0

        for row in range(self.table.rowCount()):
            try:
                name_item = self.table.item(row, 1)
                if not name_item:
                    continue

                name = name_item.text()

                is_kit_row = name.startswith("📦")
                is_child_row = name.startswith("  ├──")

                num_item = self.table.item(row, 0)
                has_number = num_item and num_item.text()

                length = self.table.item(row, 2).text() if self.table.item(row, 2) else "0"
                qty_text = self.table.item(row, 3).text() if self.table.item(row, 3) else "1"
                unit = self.table.item(row, 4).text() if self.table.item(row, 4) else "м.п."
                price_text = self.table.item(row, 5).text() if self.table.item(row, 5) else "0"
                total_text = self.table.item(row, 6).text() if self.table.item(row, 6) else "0"
                weight_text = self.table.item(row, 7).text() if self.table.item(row, 7) else "0"

                try:
                    price = float(price_text)
                except (ValueError, AttributeError):
                    price = 0.0

                try:
                    total = float(total_text)
                except (ValueError, AttributeError):
                    total = 0.0

                try:
                    weight = float(weight_text)
                except (ValueError, AttributeError):
                    weight = 0.0

                match = re.search(r'(\d+\.?\d*)', qty_text)
                qty = float(match.group(1)) if match else 1.0

                if is_child_row:
                    row_data = {
                        'num': '',
                        'name': name,
                        'length': length,
                        'quantity': qty,
                        'qty_text': qty_text,
                        'unit': unit,
                        'price': price,
                        'total': total,
                        'weight': weight,
                        'is_child': True
                    }
                elif is_kit_row:
                    row_data = {
                        'num': len([r for r in rows if not r.get('is_child')]) + 1,
                        'name': name,
                        'length': length,
                        'quantity': qty,
                        'qty_text': qty_text,
                        'unit': unit,
                        'price': price,
                        'total': total,
                        'weight': weight,
                        'is_child': False
                    }
                    total_sum += total
                    total_weight += weight
                else:
                    row_data = {
                        'num': len([r for r in rows if not r.get('is_child')]) + 1,
                        'name': name,
                        'length': length,
                        'quantity': qty,
                        'qty_text': qty_text,
                        'unit': unit,
                        'price': price,
                        'total': total,
                        'weight': weight,
                        'is_child': False
                    }
                    total_sum += total
                    total_weight += weight

                rows.append(row_data)

            except (ValueError, AttributeError) as e:
                print(f"Ошибка в строке {row}: {e}")
                continue

        return rows, total_sum, total_weight

    def _get_signer_name(self):
        if self.current_user and self.current_user.get('full_name'):
            return self.current_user['full_name']
        return "Не указан"

    def export_to_excel(self):
        if not EXCEL_AVAILABLE:
            QMessageBox.critical(self, "Ошибка", "Библиотека openpyxl не установлена!")
            return

        rows, total_sum, total_weight = self._collect_table_data()
        if not rows:
            QMessageBox.warning(self, "Ошибка", "Нет данных для экспорта!")
            return

        default_name = f"Смета_{self.client_name}_{datetime.now().strftime('%Y%m%d')}.xlsx"
        default_name = re.sub(r'[\\/*?:"<>|]', "_", default_name)

        file_path, _ = QFileDialog.getSaveFileName(self, "Сохранить смету в Excel", default_name,
                                                   "Excel файлы (*.xlsx)")
        if not file_path:
            return

        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Смета"

            title_font = Font(name='Arial', size=16, bold=True)
            header_font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
            header_fill = PatternFill(start_color='4CAF50', end_color='4CAF50', fill_type='solid')
            info_font = Font(name='Arial', size=11)
            info_bold = Font(name='Arial', size=11, bold=True)
            total_font = Font(name='Arial', size=12, bold=True, color='0055AA')
            thin_border = Border(left=Side(style='thin'), right=Side(style='thin'),
                                 top=Side(style='thin'), bottom=Side(style='thin'))
            center_align = Alignment(horizontal='center', vertical='center')
            right_align = Alignment(horizontal='right', vertical='center')

            ws.merge_cells('A1:G1')
            ws['A1'] = 'СМЕТА НА МОНТАЖ ОТКАТНЫХ ВОРОТ'
            ws['A1'].font = title_font
            ws['A1'].alignment = Alignment(horizontal='center')

            ws['A3'] = 'Номер сметы:'
            ws['A3'].font = info_bold
            ws['B3'] = f"СМ-{datetime.now().strftime('%Y%m%d')}-001"
            ws['B3'].font = info_font

            ws['A4'] = 'Дата:'
            ws['A4'].font = info_bold
            ws['B4'] = datetime.now().strftime('%d.%m.%Y')
            ws['B4'].font = info_font

            ws['A5'] = 'Клиент:'
            ws['A5'].font = info_bold
            ws['B5'] = self.client_name
            ws['B5'].font = info_font

            ws['A6'] = 'Адрес:'
            ws['A6'].font = info_bold
            ws['B6'] = self.client_address
            ws['B6'].font = info_font

            headers = ['№', 'Наименование', 'Длина/Кол-во', 'Хлысты/Ед.', 'Ед.', 'Цена (руб)', 'Сумма (руб)']
            header_row = 8

            for col_idx, header in enumerate(headers, 1):
                cell = ws.cell(row=header_row, column=col_idx, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center_align
                cell.border = thin_border

            for row_idx, item in enumerate(rows, header_row + 1):
                ws.cell(row=row_idx, column=1, value=item['num']).alignment = center_align
                ws.cell(row=row_idx, column=2, value=item['name'])
                ws.cell(row=row_idx, column=3, value=item['length']).alignment = right_align
                ws.cell(row=row_idx, column=4, value=item['qty_text']).alignment = center_align
                ws.cell(row=row_idx, column=5, value=item['unit']).alignment = center_align
                ws.cell(row=row_idx, column=6, value=item['price']).alignment = right_align
                ws.cell(row=row_idx, column=6).number_format = '#,##0.00'
                ws.cell(row=row_idx, column=7, value=item['total']).alignment = right_align
                ws.cell(row=row_idx, column=7).number_format = '#,##0.00'
                for col in range(1, 8):
                    ws.cell(row=row_idx, column=col).border = thin_border

            total_row = header_row + len(rows) + 1
            ws.merge_cells(f'A{total_row}:F{total_row}')
            ws.cell(row=total_row, column=1, value='ИТОГО:').font = total_font
            ws.cell(row=total_row, column=1).alignment = Alignment(horizontal='right')
            ws.cell(row=total_row, column=7, value=total_sum).font = total_font
            ws.cell(row=total_row, column=7).alignment = right_align
            ws.cell(row=total_row, column=7).number_format = '#,##0.00'

            weight_row = total_row + 1
            ws.merge_cells(f'A{weight_row}:F{weight_row}')
            ws.cell(row=weight_row, column=1, value='Общий вес:').font = info_bold
            ws.cell(row=weight_row, column=1).alignment = Alignment(horizontal='right')
            ws.cell(row=weight_row, column=7, value=f"{total_weight:.2f} кг").font = info_bold
            ws.cell(row=weight_row, column=7).alignment = right_align

            ws.column_dimensions['A'].width = 5
            ws.column_dimensions['B'].width = 40
            ws.column_dimensions['C'].width = 15
            ws.column_dimensions['D'].width = 28
            ws.column_dimensions['E'].width = 8
            ws.column_dimensions['F'].width = 14
            ws.column_dimensions['G'].width = 16

            signer = self._get_signer_name()
            sign_row = weight_row + 3
            ws.cell(row=sign_row, column=1, value=f'Смету составил: _________________ / {signer}').font = info_font
            ws.cell(row=sign_row + 2, column=1,
                    value=f'Документ сформирован: {datetime.now().strftime("%d.%m.%Y %H:%M")}').font = Font(
                name='Arial', size=9, italic=True, color='999999')

            wb.save(file_path)
            QMessageBox.information(self, "Готово", f"✅ Смета сохранена!\n\n{file_path}")

        except Exception as e:
            QMessageBox.critical(self, "Ошибка экспорта", f"Не удалось сохранить файл:\n\n{e}")

    def export_to_pdf(self):
        if not PDF_AVAILABLE:
            QMessageBox.critical(self, "Ошибка", "Библиотека WEasyPrint не установлена!")
            return

        rows, total_sum, total_weight = self._collect_table_data()
        if not rows:
            QMessageBox.warning(self, "Ошибка", "Нет данных для экспорта!")
            return

        default_name = f"Смета_{self.client_name}_{datetime.now().strftime('%Y%m%d')}.pdf"
        default_name = re.sub(r'[\\/*?:"<>|]', "_", default_name)

        file_path, _ = QFileDialog.getSaveFileName(self, "Сохранить смету в PDF", default_name, "PDF файлы (*.pdf)")
        if not file_path:
            return

        try:
            html_content = self._generate_pdf_html(rows, total_sum, total_weight)
            HTML(string=html_content).write_pdf(file_path)
            QMessageBox.information(self, "Готово", f"✅ PDF сохранён!\n\n{file_path}")
        except Exception as e:
            QMessageBox.critical(self, "Ошибка экспорта", f"Не удалось создать PDF:\n\n{e}")

    def _generate_pdf_html(self, rows, total_sum, total_weight):
        rows_html = ""
        for item in rows:
            rows_html += f"""
                <tr>
                    <td class="center">{item['num']}</td>
                    <td>{item['name']}</td>
                    <td class="right">{item['length']}</td>
                    <td class="center">{item['qty_text']}</td>
                    <td class="center">{item['unit']}</td>
                    <td class="right">{item['price']:,.2f}</td>
                    <td class="right"><b>{item['total']:,.2f}</b></td>
                </tr>
            """

        signer = self._get_signer_name()

        html = f"""
        <!DOCTYPE html>
        <html lang="ru">
        <head>
            <meta charset="utf-8">
            <style>
                @page {{ size: A4; margin: 2cm; }}
                body {{ font-family: 'Arial', sans-serif; font-size: 11px; color: #333; line-height: 1.4; }}
                .header {{ text-align: center; border-bottom: 3px solid #4CAF50; padding-bottom: 15px; margin-bottom: 20px; }}
                .header h1 {{ margin: 0; font-size: 22px; color: #2E7D32; }}
                .info-block {{ margin-bottom: 20px; padding: 12px; background-color: #f9f9f9; border-left: 4px solid #4CAF50; }}
                .estimate-table {{ width: 100%; border-collapse: collapse; margin: 20px 0; font-size: 10px; }}
                .estimate-table th {{ background-color: #4CAF50; color: white; padding: 8px 5px; text-align: center; border: 1px solid #3d8b40; }}
                .estimate-table td {{ padding: 6px 5px; border: 1px solid #ddd; }}
                .estimate-table tr:nth-child(even) {{ background-color: #f9f9f9; }}
                .center {{ text-align: center; }}
                .right {{ text-align: right; }}
                .totals {{ margin-top: 20px; padding: 15px; background-color: #E8F5E9; border: 2px solid #4CAF50; }}
                .grand-total {{ font-size: 18px; font-weight: bold; color: #2E7D32; }}
                .signature td {{ padding: 10px; vertical-align: top; }}
                .signature .line {{ border-bottom: 1px solid #333; display: inline-block; min-width: 200px; }}
            </style>
        </head>
        <body>
            <div class="header">
                <h1>СМЕТА</h1>
                <div>на монтаж откатных ворот</div>
            </div>

            <div class="info-block">
                <table>
                    <tr><td><b>Номер сметы:</b></td><td>СМ-{datetime.now().strftime('%Y%m%d')}-001</td>
                        <td><b>Дата:</b></td><td>{datetime.now().strftime('%d.%m.%Y')}</td></tr>
                    <tr><td><b>Клиент:</b></td><td colspan="3"><b>{self.client_name}</b></td></tr>
                    <tr><td><b>Адрес:</b></td><td colspan="3">{self.client_address}</td></tr>
                </table>
            </div>

            <table class="estimate-table">
                <thead>
                    <tr>
                        <th style="width: 4%;">№</th>
                        <th style="width: 34%;">Наименование</th>
                        <th style="width: 10%;">Длина/Кол-во</th>
                        <th style="width: 16%;">Хлысты/Ед.</th>
                        <th style="width: 8%;">Ед.</th>
                        <th style="width: 14%;">Цена, ₽</th>
                        <th style="width: 14%;">Сумма, ₽</th>
                    </tr>
                </thead>
                <tbody>{rows_html}</tbody>
            </table>

            <div class="totals">
                <div><b>Общий вес материалов:</b> {total_weight:,.2f} кг</div>
                <div class="grand-total">ИТОГО К ОПЛАТЕ: {total_sum:,.2f} ₽</div>
            </div>

            <div class="signature">
                <table style="width: 100%;">
                    <tr>
                        <td>Смету составил:<br><span class="line"></span> / {signer}</td>
                        <td style="text-align: right;">Заказчик:<br><span class="line"></span> / {self.client_name}</td>
                    </tr>
                </table>
            </div>

            <div style="margin-top: 30px; font-size: 9px; color: #888; text-align: center;">
                Документ сформирован автоматически {datetime.now().strftime('%d.%m.%Y в %H:%M')}
            </div>
        </body>
        </html>
        """
        return html

    def print_pdf(self):
        rows = []
        total_sum = 0.0
        for row in range(self.table.rowCount()):
            try:
                num_item = self.table.item(row, 0)
                if not num_item or not num_item.text():
                    continue

                name = self.table.item(row, 1).text()
                qty_text = self.table.item(row, 3).text()
                match = re.search(r'(\d+)', qty_text)
                qty = float(match.group(1)) if match else 1.0
                unit = self.table.item(row, 4).text()
                price = float(self.table.item(row, 5).text())
                total = float(self.table.item(row, 6).text())
                rows.append((name, qty, unit, price, total))
                total_sum += total
            except (ValueError, AttributeError):
                pass

        html = self.generate_pdf_html(rows, total_sum)

        printer = QPrinter(QPrinter.PrinterMode.HighResolution)
        printer.setOutputFormat(QPrinter.OutputFormat.PdfFormat)
        printer.setOutputFileName(f"Смета_{datetime.now().strftime('%Y%m%d')}.pdf")

        dialog = QPrintDialog(printer, self)
        if dialog.exec() == QPrintDialog.DialogCode.Accepted:
            doc = QTextDocument()
            doc.setHtml(html)
            doc.print_(printer)

    def generate_pdf_html(self, rows, total_sum):
        items_html = ""
        for i, (name, qty, unit, price, total) in enumerate(rows, 1):
            items_html += f"""
            <tr>
                <td style="border:1px solid #ddd;padding:4px;text-align:center;">{i}</td>
                <td style="border:1px solid #ddd;padding:4px;">{name}</td>
                <td style="border:1px solid #ddd;padding:4px;text-align:center;">{qty:.0f}</td>
                <td style="border:1px solid #ddd;padding:4px;text-align:center;">{unit}</td>
                <td style="border:1px solid #ddd;padding:4px;text-align:right;">{price:.2f}</td>
                <td style="border:1px solid #ddd;padding:4px;text-align:right;">{total:.2f}</td>
            </tr>
            """

        signer = self._get_signer_name()

        return f"""
        <html>
        <head>
            <meta charset="utf-8">
            <style>
                body {{ font-family: Arial, sans-serif; margin: 40px; }}
                h1 {{ text-align: center; font-size: 18px; }}
                table {{ width: 100%; border-collapse: collapse; margin-top: 20px; }}
                th {{ background-color: #4CAF50; color: white; padding: 8px; border: 1px solid #ddd; }}
                td {{ border: 1px solid #ddd; padding: 4px; }}
                .total {{ font-weight: bold; font-size: 16px; text-align: right; margin-top: 20px; }}
            </style>
        </head>
        <body>
            <h1>СМЕТА</h1>
            <table>
                <tr><td><b>Клиент:</b> {self.client_name}</td><td><b>Дата:</b> {datetime.now().strftime('%d.%m.%Y')}</td></tr>
                <tr><td><b>Адрес:</b> {self.client_address}</td><td></td></tr>
            </table>
            <table>
                <thead>
                    <tr>
                        <th style="width:30px;">№</th>
                        <th>Наименование</th>
                        <th style="width:60px;">Кол-во</th>
                        <th style="width:50px;">Ед.</th>
                        <th style="width:80px;">Цена</th>
                        <th style="width:90px;">Сумма</th>
                    </tr>
                </thead>
                <tbody>{items_html}</tbody>
            </table>
            <div class="total">ИТОГО: {total_sum:.2f} руб</div>
            <div style="margin-top: 30px;">Смету составил: _________________ / {signer}</div>
        </body>
        </html>
        """


class SelectMaterialTreeDialog(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Выбор из номенклатуры")
        self.setMinimumSize(750, 550)
        self.selected_data = None

        layout = QVBoxLayout(self)

        hint = QLabel(
            "📁 Жёлтые папки — категории\n"
            "📄 Белые — материалы\n"
            "📦 Зелёные — комплекты\n\n"
            "💡 Кликните на треугольник слева от папки, чтобы раскрыть её"
        )
        hint.setStyleSheet("background-color: #f0f0f0; padding: 8px; border-radius: 4px;")
        hint.setWordWrap(True)
        layout.addWidget(hint)

        search_layout = QHBoxLayout()
        search_layout.addWidget(QLabel("🔍 Поиск:"))
        self.search_input = QLineEdit()
        self.search_input.setPlaceholderText("Введите название...")
        self.search_input.textChanged.connect(self.filter_items)
        search_layout.addWidget(self.search_input)
        layout.addLayout(search_layout)

        self.tree = QTreeWidget()
        self.tree.setHeaderLabels(["Наименование", "Ед. изм.", "Цена"])
        self.tree.setColumnWidth(0, 350)
        self.tree.setColumnWidth(1, 80)
        self.tree.setColumnWidth(2, 100)
        self.tree.itemDoubleClicked.connect(self.select_item)
        layout.addWidget(self.tree)

        btn_layout = QHBoxLayout()
        self.select_btn = QPushButton("✅ Выбрать")
        self.select_btn.clicked.connect(self.select_item)
        btn_layout.addWidget(self.select_btn)

        self.cancel_btn = QPushButton("❌ Отмена")
        self.cancel_btn.clicked.connect(self.reject)
        btn_layout.addWidget(self.cancel_btn)

        btn_layout.addStretch()
        layout.addLayout(btn_layout)

        self.all_materials = []
        self.all_kits = []
        self.all_categories = []
        self.all_items_data = {}
        self._items_dict = {}

        self.load_data()

    def _get_icon_for_type(self, item_type, is_kit=False):
        pixmap = QPixmap(20, 20)
        pixmap.fill(Qt.GlobalColor.transparent)
        painter = QPainter(pixmap)
        painter.setRenderHint(QPainter.RenderHint.Antialiasing)

        if item_type == 'category':
            painter.setBrush(QBrush(QColor(255, 193, 7)))
            painter.setPen(QPen(QColor(255, 160, 0), 1))
            painter.drawRect(2, 5, 16, 12)
            painter.drawLine(2, 5, 6, 2)
            painter.drawLine(6, 2, 18, 2)
        elif is_kit:
            painter.setBrush(QBrush(QColor(76, 175, 80)))
            painter.setPen(QPen(QColor(46, 125, 50), 1))
            painter.drawRect(3, 4, 14, 12)
            painter.drawLine(3, 8, 17, 8)
            painter.drawArc(7, 2, 6, 4, 0, 180 * 16)
        else:
            painter.setBrush(QBrush(QColor(255, 255, 255)))
            painter.setPen(QPen(QColor(100, 100, 100), 1))
            painter.drawRect(3, 3, 14, 14)
            painter.drawLine(5, 8, 15, 8)
            painter.drawLine(5, 12, 13, 12)

        painter.end()
        return QIcon(pixmap)

    def load_data(self):
        self.tree.clear()
        self._items_dict = {}

        self.all_materials = db.get_materials_hierarchy()
        self.all_kits = db.get_all_kits()
        self.all_categories = [m for m in self.all_materials if m.get('is_category', False)]

        categories = [m for m in self.all_materials if m.get('is_category', False)]
        regular_materials = [m for m in self.all_materials if not m.get('is_category', False)]

        categories.sort(key=lambda x: x['name'])
        regular_materials.sort(key=lambda x: x['name'])

        for cat in categories:
            item = QTreeWidgetItem()
            item.setText(0, cat['name'])
            item.setText(1, 'категория')
            item.setText(2, '')
            item.setData(0, Qt.ItemDataRole.UserRole, {
                'type': 'category',
                'id': cat['id'],
                'name': cat['name']
            })
            item.setIcon(0, self._get_icon_for_type('category'))
            item.setFlags(item.flags() & ~Qt.ItemFlag.ItemIsSelectable)
            self._items_dict[cat['id']] = item

        for mat in regular_materials:
            item = QTreeWidgetItem()
            item.setText(0, mat['name'])
            item.setText(1, mat.get('unit', ''))
            item.setText(2, f"{float(mat.get('retail_price', 0)):.2f}")
            item.setData(0, Qt.ItemDataRole.UserRole, {
                'type': 'material',
                'id': mat['id'],
                'name': mat['name'],
                'unit': mat.get('unit', 'м.п.'),
                'retail_price': float(mat.get('retail_price', 0))
            })
            item.setIcon(0, self._get_icon_for_type('material'))
            self._items_dict[mat['id']] = item

        for kit in self.all_kits:
            kit_item = QTreeWidgetItem()
            kit_item.setText(0, kit['name'])
            kit_item.setText(1, 'комплект')
            kit_item.setText(2, f"{float(kit.get('retail_price', 0)):.2f}")
            kit_item.setData(0, Qt.ItemDataRole.UserRole, {
                'type': 'kit',
                'id': kit['id'],
                'name': kit['name']
            })
            kit_item.setIcon(0, self._get_icon_for_type('kit', is_kit=True))
            self._items_dict[f"kit_{kit['id']}"] = kit_item

        for cat in categories:
            if cat['parent_id'] is None:
                self.tree.addTopLevelItem(self._items_dict[cat['id']])
            else:
                parent_item = self._items_dict.get(cat['parent_id'])
                if parent_item:
                    parent_item.addChild(self._items_dict[cat['id']])

        for mat in regular_materials:
            if mat['parent_id'] is None:
                self.tree.addTopLevelItem(self._items_dict[mat['id']])
            else:
                parent_item = self._items_dict.get(mat['parent_id'])
                if parent_item:
                    parent_item.addChild(self._items_dict[mat['id']])

        kit_category_id = None
        for cat in categories:
            if cat['name'] == 'Комплекты':
                kit_category_id = cat['id']
                break

        if kit_category_id and kit_category_id in self._items_dict:
            parent_item = self._items_dict[kit_category_id]
            for kit in self.all_kits:
                kit_item = self._items_dict.get(f"kit_{kit['id']}")
                if kit_item:
                    parent_item.addChild(kit_item)
        else:
            for kit in self.all_kits:
                kit_item = self._items_dict.get(f"kit_{kit['id']}")
                if kit_item:
                    self.tree.addTopLevelItem(kit_item)

        self.tree.expandAll()
        self.tree.resizeColumnToContents(0)

    def filter_items(self):
        query = self.search_input.text().strip().lower()

        self.load_data()

        if not query:
            return

        all_items = []
        for i in range(self.tree.topLevelItemCount()):
            all_items.append(self.tree.topLevelItem(i))
            for j in range(self.tree.topLevelItem(i).childCount()):
                all_items.append(self.tree.topLevelItem(i).child(j))

        for item in all_items:
            item.setHidden(True)

        for item in all_items:
            data = item.data(0, Qt.ItemDataRole.UserRole)
            if data:
                name = data.get('name', '').lower()
                if query in name:
                    item.setHidden(False)
                    parent = item.parent()
                    if parent:
                        parent.setHidden(False)

        self.tree.expandAll()

    def get_selected_item(self):
        items = self.tree.selectedItems()
        if not items:
            return None
        return items[0]

    def select_item(self):
        item = self.get_selected_item()
        if not item:
            QMessageBox.warning(self, "Ошибка", "Выберите элемент!")
            return

        data = item.data(0, Qt.ItemDataRole.UserRole)
        if not data:
            return

        if data.get('type') == 'category':
            QMessageBox.warning(self, "Ошибка", "Нельзя выбрать категорию. Выберите материал или комплект!")
            return

        self.selected_data = data
        self.accept()

    def get_selected_data(self):
        return self.selected_data