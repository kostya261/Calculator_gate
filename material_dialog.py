import json
from PyQt6.QtWidgets import *
from PyQt6.QtCore import *
from PyQt6.QtGui import QColor
import db


class MaterialDialog(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Справочник материалов")
        self.setMinimumSize(950, 500)

        layout = QVBoxLayout(self)

        # ===== Таблица =====
        self.table = QTableWidget()
        self.table.setColumnCount(8)
        self.table.setHorizontalHeaderLabels(
            ["ID", "Артикул", "Наименование", "Ед. изм.", "Вес", "Цена закуп.", "Цена розн.", "Описание"]
        )
        self.table.horizontalHeader().setStretchLastSection(False)
        self.table.setColumnWidth(0, 40)
        self.table.setColumnWidth(1, 80)
        self.table.setColumnWidth(2, 200)
        self.table.setColumnWidth(3, 70)
        self.table.setColumnWidth(4, 70)
        self.table.setColumnWidth(5, 90)
        self.table.setColumnWidth(6, 90)
        self.table.horizontalHeader().setSectionResizeMode(7, QHeaderView.ResizeMode.Stretch)

        self.table.cellChanged.connect(self.on_cell_changed)

        layout.addWidget(self.table)

        hint_label = QLabel("💡 Совет: можно менять цены прямо в таблице (колонки 'Цена закуп.' и 'Цена розн.')")
        hint_label.setStyleSheet("color: #666; font-style: italic; padding: 5px;")
        layout.addWidget(hint_label)

        btn_layout = QHBoxLayout()

        self.add_btn = QPushButton("➕ Добавить")
        self.add_btn.clicked.connect(self.add_material)
        btn_layout.addWidget(self.add_btn)

        self.edit_btn = QPushButton("✏️ Редактировать")
        self.edit_btn.clicked.connect(self.edit_material)
        btn_layout.addWidget(self.edit_btn)

        self.delete_btn = QPushButton("🗑️ Удалить")
        self.delete_btn.clicked.connect(self.delete_material)
        btn_layout.addWidget(self.delete_btn)

        self.refresh_btn = QPushButton("🔄 Обновить")
        self.refresh_btn.clicked.connect(self.load_data)
        btn_layout.addWidget(self.refresh_btn)

        btn_layout.addStretch()

        self.import_excel_btn = QPushButton("📥 Импорт Excel")
        self.import_excel_btn.clicked.connect(self.import_excel)
        btn_layout.addWidget(self.import_excel_btn)

        self.export_excel_btn = QPushButton("📤 Экспорт Excel")
        self.export_excel_btn.clicked.connect(self.export_excel)
        btn_layout.addWidget(self.export_excel_btn)

        self.import_json_btn = QPushButton("📥 Импорт JSON")
        self.import_json_btn.clicked.connect(self.import_json)
        btn_layout.addWidget(self.import_json_btn)

        self.export_json_btn = QPushButton("📤 Экспорт JSON")
        self.export_json_btn.clicked.connect(self.export_json)
        btn_layout.addWidget(self.export_json_btn)

        layout.addLayout(btn_layout)

        self.load_data()

    def load_data(self):
        self.table.blockSignals(True)
        try:
            materials = db.get_materials()
            self.table.setRowCount(len(materials))

            for row, mat in enumerate(materials):
                id_item = QTableWidgetItem(str(mat['id']))
                id_item.setFlags(id_item.flags() & ~Qt.ItemFlag.ItemIsEditable)
                self.table.setItem(row, 0, id_item)

                sku_item = QTableWidgetItem(mat['sku'] or '')
                sku_item.setFlags(sku_item.flags() & ~Qt.ItemFlag.ItemIsEditable)
                self.table.setItem(row, 1, sku_item)

                name_item = QTableWidgetItem(mat['name'])
                name_item.setFlags(name_item.flags() & ~Qt.ItemFlag.ItemIsEditable)
                self.table.setItem(row, 2, name_item)

                unit_item = QTableWidgetItem(mat['unit'])
                unit_item.setFlags(unit_item.flags() & ~Qt.ItemFlag.ItemIsEditable)
                self.table.setItem(row, 3, unit_item)

                weight = mat['weight_per_unit'] if mat['weight_per_unit'] is not None else 0
                weight_item = QTableWidgetItem(f"{weight:.3f}")
                weight_item.setFlags(weight_item.flags() & ~Qt.ItemFlag.ItemIsEditable)
                self.table.setItem(row, 4, weight_item)

                purchase = mat['purchase_price'] if mat['purchase_price'] is not None else 0
                purchase_item = QTableWidgetItem(f"{purchase:.2f}")
                self.table.setItem(row, 5, purchase_item)

                retail = mat['retail_price'] if mat['retail_price'] is not None else 0
                retail_item = QTableWidgetItem(f"{retail:.2f}")
                self.table.setItem(row, 6, retail_item)

                desc_item = QTableWidgetItem(mat['description'] or '')
                desc_item.setFlags(desc_item.flags() & ~Qt.ItemFlag.ItemIsEditable)
                self.table.setItem(row, 7, desc_item)

            self.table.resizeRowsToContents()
        finally:
            self.table.blockSignals(False)

    def on_cell_changed(self, row, col):
        if col not in [5, 6]:
            return

        try:
            id_item = self.table.item(row, 0)
            if not id_item:
                return
            material_id = int(id_item.text())

            price_item = self.table.item(row, col)
            if not price_item:
                return

            price_text = price_item.text().strip()
            if not price_text or price_text.lower() == 'none':
                new_price = 0.0
            else:
                new_price = float(price_text)

            material = db.get_material_by_id(material_id)
            if not material:
                return

            if col == 5:
                purchase_price = new_price
                retail_price = material['retail_price'] if material['retail_price'] else 0
            else:
                purchase_price = material['purchase_price'] if material['purchase_price'] else 0
                retail_price = new_price

            self.table.blockSignals(True)
            try:
                success = db.update_material(
                    material_id,
                    material['name'],
                    material['unit'],
                    material['weight_per_unit'] if material['weight_per_unit'] else 0,
                    purchase_price,
                    retail_price,
                    material['sku'],
                    material['description'] or ''
                )

                if success:
                    price_item.setBackground(QColor(200, 255, 200))
                    QTimer.singleShot(1000, lambda r=row, c=col: self._reset_cell_color(r, c))
                else:
                    QMessageBox.warning(self, "Ошибка", "Не удалось сохранить цену!")
                    old_price = material['purchase_price'] if col == 5 else material['retail_price']
                    price_item.setText(f"{old_price:.2f}" if old_price else "0.00")
            finally:
                self.table.blockSignals(False)

        except ValueError:
            QMessageBox.warning(self, "Ошибка", "Цена должна быть числом!")
            self.table.blockSignals(True)
            try:
                material_id = int(self.table.item(row, 0).text())
                material = db.get_material_by_id(material_id)
                old_price = material['purchase_price'] if col == 5 else material['retail_price']
                price_item.setText(f"{old_price:.2f}" if old_price else "0.00")
            finally:
                self.table.blockSignals(False)
        except Exception as e:
            print(f"Ошибка при сохранении цены: {e}")

    def _reset_cell_color(self, row, col):
        self.table.blockSignals(True)
        try:
            item = self.table.item(row, col)
            if item:
                item.setBackground(QColor(255, 255, 255))
        finally:
            self.table.blockSignals(False)

    def get_selected_id(self):
        row = self.table.currentRow()
        if row < 0:
            return None
        id_item = self.table.item(row, 0)
        return int(id_item.text()) if id_item else None

    def add_material(self):
        dialog = MaterialEditDialog(self)
        if dialog.exec() == QDialog.DialogCode.Accepted:
            self.load_data()

    def edit_material(self):
        material_id = self.get_selected_id()
        if material_id is None:
            QMessageBox.warning(self, "Ошибка", "Выберите материал для редактирования!")
            return

        dialog = MaterialEditDialog(self, material_id)
        if dialog.exec() == QDialog.DialogCode.Accepted:
            self.load_data()

    def delete_material(self):
        material_id = self.get_selected_id()
        if material_id is None:
            QMessageBox.warning(self, "Ошибка", "Выберите материал для удаления!")
            return

        reply = QMessageBox.question(
            self,
            "Подтверждение",
            "Удалить выбранный материал?",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
        )

        if reply == QMessageBox.StandardButton.Yes:
            if db.delete_material(material_id):
                self.load_data()
            else:
                QMessageBox.warning(self, "Ошибка", "Не удалось удалить материал!")

    def import_excel(self):
        file_path, _ = QFileDialog.getOpenFileName(
            self,
            "Выберите Excel файл для импорта",
            "",
            "Excel files (*.xlsx *.xls)"
        )
        if file_path:
            try:
                from openpyxl import load_workbook
                wb = load_workbook(file_path)
                ws = wb.active

                imported = 0
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if not row[0]:
                        continue
                    name = row[0]
                    unit = row[1] or 'м.п.'
                    weight = float(row[2]) if row[2] else 0
                    purchase_price = float(row[3]) if row[3] else 0
                    retail_price = float(row[4]) if row[4] else 0
                    sku = str(row[5]) if row[5] else None
                    description = str(row[6]) if row[6] else ''

                    db.add_material(name, unit, weight, purchase_price, retail_price, sku, description)
                    imported += 1

                QMessageBox.information(self, "Готово", f"Импортировано {imported} материалов.")
                self.load_data()
            except Exception as e:
                QMessageBox.warning(self, "Ошибка импорта", str(e))

    def export_excel(self):
        file_path, _ = QFileDialog.getSaveFileName(
            self,
            "Сохранить Excel файл",
            "materials.xlsx",
            "Excel files (*.xlsx)"
        )
        if file_path:
            try:
                from openpyxl import Workbook
                wb = Workbook()
                ws = wb.active
                ws.title = "Материалы"

                headers = ["Наименование", "Ед. изм.", "Вес", "Цена закуп.", "Цена розн.", "Артикул", "Описание"]
                ws.append(headers)

                materials = db.get_materials()
                for mat in materials:
                    ws.append([
                        mat['name'],
                        mat['unit'],
                        mat['weight_per_unit'],
                        mat['purchase_price'],
                        mat['retail_price'],
                        mat['sku'],
                        mat['description']
                    ])

                wb.save(file_path)
                QMessageBox.information(self, "Готово", f"Экспортировано {len(materials)} материалов.")
            except Exception as e:
                QMessageBox.warning(self, "Ошибка экспорта", str(e))

    def import_json(self):
        file_path, _ = QFileDialog.getOpenFileName(
            self,
            "Выберите JSON файл для импорта",
            "",
            "JSON files (*.json)"
        )
        if file_path:
            try:
                with open(file_path, 'r', encoding='utf-8') as f:
                    data = json.load(f)

                imported = 0
                for item in data:
                    db.add_material(
                        name=item.get('name'),
                        unit=item.get('unit', 'м.п.'),
                        weight=item.get('weight_per_unit', 0),
                        purchase_price=item.get('purchase_price', 0),
                        retail_price=item.get('retail_price', 0),
                        sku=item.get('sku'),
                        description=item.get('description', '')
                    )
                    imported += 1

                QMessageBox.information(self, "Готово", f"Импортировано {imported} материалов.")
                self.load_data()
            except Exception as e:
                QMessageBox.warning(self, "Ошибка импорта", str(e))

    def export_json(self):
        file_path, _ = QFileDialog.getSaveFileName(
            self,
            "Сохранить JSON файл",
            "materials.json",
            "JSON files (*.json)"
        )
        if file_path:
            try:
                materials = db.get_materials()
                with open(file_path, 'w', encoding='utf-8') as f:
                    json.dump(materials, f, ensure_ascii=False, indent=2, default=str)

                QMessageBox.information(self, "Готово", f"Экспортировано {len(materials)} материалов.")
            except Exception as e:
                QMessageBox.warning(self, "Ошибка экспорта", str(e))


class MaterialEditDialog(QDialog):
    def __init__(self, parent=None, material_id=None, parent_id=None):
        super().__init__(parent)
        self.material_id = material_id
        self.parent_id = parent_id
        self.setWindowTitle("Редактирование материала" if material_id else "Новый материал")
        self.setMinimumWidth(480)

        layout = QVBoxLayout(self)
        form_layout = QFormLayout()

        self.name_edit = QLineEdit()
        form_layout.addRow("Наименование:", self.name_edit)

        self.sku_edit = QLineEdit()
        form_layout.addRow("Артикул:", self.sku_edit)

        self.unit_combo = QComboBox()
        self.load_units()
        form_layout.addRow("Единица измерения:", self.unit_combo)

        self.weight_edit = QLineEdit()
        self.weight_edit.setPlaceholderText("0.00")
        form_layout.addRow("Вес (кг):", self.weight_edit)

        self.purchase_price_edit = QLineEdit()
        self.purchase_price_edit.setPlaceholderText("0.00")
        form_layout.addRow("Закупочная цена:", self.purchase_price_edit)

        self.retail_price_edit = QLineEdit()
        self.retail_price_edit.setPlaceholderText("0.00")
        form_layout.addRow("Розничная цена:", self.retail_price_edit)

        # ⚠️ КАТЕГОРИЯ С КНОПКОЙ ...
        category_layout = QHBoxLayout()
        self.parent_combo = QComboBox()
        self.parent_combo.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        category_layout.addWidget(self.parent_combo)

        self.category_btn = QPushButton("...")
        self.category_btn.setFixedWidth(30)
        self.category_btn.setToolTip("Выбрать категорию из дерева")
        self.category_btn.clicked.connect(self.select_category_from_tree)
        category_layout.addWidget(self.category_btn)

        form_layout.addRow("Категория:", category_layout)

        self.desc_edit = QTextEdit()
        self.desc_edit.setMaximumHeight(80)
        form_layout.addRow("Описание:", self.desc_edit)

        layout.addLayout(form_layout)

        btn_layout = QHBoxLayout()
        self.save_btn = QPushButton("💾 Сохранить")
        self.save_btn.clicked.connect(self.save)
        btn_layout.addWidget(self.save_btn)

        self.cancel_btn = QPushButton("❌ Отмена")
        self.cancel_btn.clicked.connect(self.reject)
        btn_layout.addWidget(self.cancel_btn)

        layout.addLayout(btn_layout)

        self.load_categories()
        if material_id:
            self.load_data()

    def load_units(self):
        units = db.get_units()
        self.unit_combo.clear()
        for unit in units:
            self.unit_combo.addItem(unit['name'], unit['id'])

    def load_categories(self):
        """Загружает все категории в выпадающий список с правильными отступами"""
        self.parent_combo.clear()
        self.parent_combo.addItem("(Корневая категория)", None)

        materials = db.get_materials_hierarchy()
        categories = [m for m in materials if m.get('is_category', False)]

        # Находим корневые категории (parent_id IS NULL)
        root_cats = [cat for cat in categories if cat['parent_id'] is None]
        root_cats.sort(key=lambda x: x['name'])

        # Рекурсивно добавляем категории с отступами
        def add_category_with_indent(cat, indent=0):
            display_name = "  " * indent + f"📁 {cat['name']}"
            self.parent_combo.addItem(display_name, cat['id'])

            # Находим дочерние категории
            children = [c for c in categories if c['parent_id'] == cat['id']]
            children.sort(key=lambda x: x['name'])
            for child in children:
                add_category_with_indent(child, indent + 1)

        # Добавляем все корневые категории
        for root_cat in root_cats:
            add_category_with_indent(root_cat)

    def _get_category_depth(self, category_id, categories):
        """Больше не нужен, но оставляем для совместимости"""
        return 0

    def load_data(self):
        material = db.get_material_by_id(self.material_id)
        if material:
            self.name_edit.setText(material['name'] or '')
            self.sku_edit.setText(material['sku'] or '')

            units = db.get_units()
            for i, unit in enumerate(units):
                if unit['name'] == material['unit']:
                    self.unit_combo.setCurrentIndex(i)
                    break

            weight = material['weight_per_unit'] if material['weight_per_unit'] is not None else 0
            purchase = material['purchase_price'] if material['purchase_price'] is not None else 0
            retail = material['retail_price'] if material['retail_price'] is not None else 0

            self.weight_edit.setText(f"{weight:.3f}")
            self.purchase_price_edit.setText(f"{purchase:.2f}")
            self.retail_price_edit.setText(f"{retail:.2f}")
            self.desc_edit.setText(material['description'] or '')

            # Устанавливаем родительскую категорию
            parent_id = material.get('parent_id')
            if parent_id is not None:
                index = self.parent_combo.findData(parent_id)
                if index >= 0:
                    self.parent_combo.setCurrentIndex(index)
                else:
                    self.parent_combo.setCurrentIndex(0)
            else:
                self.parent_combo.setCurrentIndex(0)

    def select_category_from_tree(self):
        """Открывает диалог выбора категории с деревом"""
        current_id = self.parent_combo.currentData()
        dialog = SelectCategoryDialog(self, current_id)
        if dialog.exec() == QDialog.DialogCode.Accepted:
            selected_id = dialog.get_selected_id()
            if selected_id is not None:
                index = self.parent_combo.findData(selected_id)
                if index >= 0:
                    self.parent_combo.setCurrentIndex(index)
            else:
                # Корневая категория
                self.parent_combo.setCurrentIndex(0)

    def save(self):
        name = self.name_edit.text().strip()
        if not name:
            QMessageBox.warning(self, "Ошибка", "Наименование обязательно!")
            return

        def safe_float(text, field_name):
            text = text.strip()
            if not text or text.lower() == 'none':
                return 0.0
            try:
                return float(text)
            except ValueError:
                QMessageBox.warning(self, "Ошибка", f"{field_name} должно быть числом!")
                return None

        weight = safe_float(self.weight_edit.text(), "Вес")
        if weight is None:
            return

        purchase_price = safe_float(self.purchase_price_edit.text(), "Закупочная цена")
        if purchase_price is None:
            return

        retail_price = safe_float(self.retail_price_edit.text(), "Розничная цена")
        if retail_price is None:
            return

        sku = self.sku_edit.text().strip() or None
        unit = self.unit_combo.currentText()
        description = self.desc_edit.toPlainText().strip()

        parent_id = self.parent_combo.currentData()

        if self.material_id:
            success = db.update_material(
                self.material_id, name, unit, weight, purchase_price, retail_price, sku, description, parent_id
            )
        else:
            success = db.add_material(
                name, unit, weight, purchase_price, retail_price, sku, description, parent_id
            )

        if success:
            self.accept()
        else:
            QMessageBox.warning(self, "Ошибка", "Не удалось сохранить материал!")


class SelectCategoryDialog(QDialog):
    """Диалог выбора категории с деревом и поиском"""
    def __init__(self, parent=None, selected_id=None):
        super().__init__(parent)
        self.setWindowTitle("Выбор категории")
        self.setMinimumSize(500, 450)
        self.selected_id = selected_id
        self.result_id = None
        self.all_categories = []

        layout = QVBoxLayout(self)

        hint = QLabel("Выберите категорию для материала. Можно оставить пустым (корневая категория).")
        hint.setStyleSheet("background-color: #f0f0f0; padding: 8px; border-radius: 4px;")
        hint.setWordWrap(True)
        layout.addWidget(hint)

        # ===== СТРОКА ПОИСКА =====
        search_layout = QHBoxLayout()
        search_layout.addWidget(QLabel("🔍 Поиск:"))
        self.search_input = QLineEdit()
        self.search_input.setPlaceholderText("Введите название категории...")
        self.search_input.textChanged.connect(self.filter_categories)
        search_layout.addWidget(self.search_input)
        layout.addLayout(search_layout)

        # ===== ДЕРЕВО КАТЕГОРИЙ =====
        self.tree = QTreeWidget()
        self.tree.setHeaderLabels(["Категории"])
        self.tree.setColumnWidth(0, 400)
        self.tree.itemDoubleClicked.connect(self.select_item)
        layout.addWidget(self.tree)

        btn_layout = QHBoxLayout()

        self.select_btn = QPushButton("✅ Выбрать")
        self.select_btn.clicked.connect(self.select_item)
        btn_layout.addWidget(self.select_btn)

        self.clear_btn = QPushButton("🗑️ Очистить (корневая)")
        self.clear_btn.clicked.connect(self.clear_selection)
        btn_layout.addWidget(self.clear_btn)

        self.cancel_btn = QPushButton("❌ Отмена")
        self.cancel_btn.clicked.connect(self.reject)
        btn_layout.addWidget(self.cancel_btn)

        btn_layout.addStretch()
        layout.addLayout(btn_layout)

        self.load_data()

    def load_data(self):
        """Загружает все категории в дерево"""
        self.tree.clear()

        materials = db.get_materials_hierarchy()
        self.all_categories = [m for m in materials if m.get('is_category', False)]
        self.all_categories.sort(key=lambda x: x['name'])

        self._build_tree(self.all_categories)

    def _build_tree(self, categories):
        """Строит дерево из переданного списка категорий"""
        self.tree.clear()

        items_dict = {}

        # Создаём элементы для всех категорий
        for cat in categories:
            item = QTreeWidgetItem()
            item.setText(0, f"📁 {cat['name']}")
            item.setData(0, Qt.ItemDataRole.UserRole, cat['id'])
            items_dict[cat['id']] = item

            # Если это выбранная категория — выделяем её
            if self.selected_id == cat['id']:
                item.setSelected(True)

        # Прикрепляем к родителям
        for cat in categories:
            if cat['parent_id'] is None:
                if cat['id'] in items_dict:
                    self.tree.addTopLevelItem(items_dict[cat['id']])
            else:
                parent_item = items_dict.get(cat['parent_id'])
                if parent_item and cat['id'] in items_dict:
                    parent_item.addChild(items_dict[cat['id']])

        self.tree.expandAll()
        self.tree.resizeColumnToContents(0)

    def filter_categories(self):
        """Фильтрует категории по поиску"""
        query = self.search_input.text().strip().lower()

        if not query:
            # Показываем все категории
            self._build_tree(self.all_categories)
            return

        # Фильтруем категории, которые содержат запрос
        filtered = []
        for cat in self.all_categories:
            if query in cat['name'].lower():
                filtered.append(cat)

        # Также добавляем родителей отфильтрованных категорий
        # (чтобы сохранить иерархию)
        parent_ids = set()
        for cat in filtered:
            if cat['parent_id'] is not None:
                parent_ids.add(cat['parent_id'])

        # Добавляем всех родителей в результат
        for cat in self.all_categories:
            if cat['id'] in parent_ids and cat not in filtered:
                filtered.append(cat)

        self._build_tree(filtered)

    def select_item(self):
        """Выбирает категорию"""
        item = self.tree.currentItem()
        if not item:
            QMessageBox.warning(self, "Ошибка", "Выберите категорию!")
            return

        self.result_id = item.data(0, Qt.ItemDataRole.UserRole)
        self.accept()

    def clear_selection(self):
        """Очищает выбор (корневая категория)"""
        self.result_id = None
        self.accept()

    def get_selected_id(self):
        return self.result_id