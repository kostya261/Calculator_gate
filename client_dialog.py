import json
from PyQt6.QtWidgets import *
from PyQt6.QtCore import *
import db
from address_dialog import AddressDialog


class ClientDialog(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Справочник клиентов")
        self.setMinimumSize(1050, 500)

        layout = QVBoxLayout(self)

        self.table = QTableWidget()
        self.table.setColumnCount(10)
        self.table.setHorizontalHeaderLabels(
            ["ID", "Тип", "Фамилия", "Имя", "Отчество", "Организация",
             "Телефон 1", "Телефон 2", "Email", "Мессенджер"]
        )
        self.table.horizontalHeader().setStretchLastSection(False)
        self.table.setColumnWidth(0, 30)
        self.table.setColumnWidth(1, 60)
        self.table.setColumnWidth(2, 120)
        self.table.setColumnWidth(3, 100)
        self.table.setColumnWidth(4, 100)
        self.table.setColumnWidth(5, 150)
        self.table.setColumnWidth(6, 100)
        self.table.setColumnWidth(7, 100)
        self.table.setColumnWidth(8, 150)
        self.table.horizontalHeader().setSectionResizeMode(9, QHeaderView.ResizeMode.Stretch)
        layout.addWidget(self.table)

        btn_layout = QHBoxLayout()
        self.add_btn = QPushButton("➕ Добавить")
        self.add_btn.clicked.connect(self.add_client)
        btn_layout.addWidget(self.add_btn)

        self.edit_btn = QPushButton("✏️ Редактировать")
        self.edit_btn.clicked.connect(self.edit_client)
        btn_layout.addWidget(self.edit_btn)

        self.delete_btn = QPushButton("🗑️ Удалить")
        self.delete_btn.clicked.connect(self.delete_client)
        btn_layout.addWidget(self.delete_btn)

        self.refresh_btn = QPushButton("🔄 Обновить")
        self.refresh_btn.clicked.connect(self.load_data)
        btn_layout.addWidget(self.refresh_btn)

        self.addr_btn = QPushButton("🏠 Адреса")
        self.addr_btn.clicked.connect(self.open_addresses)
        btn_layout.addWidget(self.addr_btn)

        btn_layout.addStretch()

        self.import_excel_btn = QPushButton("📥 Импорт Excel")
        self.import_excel_btn.clicked.connect(self.import_excel)
        btn_layout.addWidget(self.import_excel_btn)

        self.export_excel_btn = QPushButton("📤 Экспорт Excel")
        self.export_excel_btn.clicked.connect(self.export_excel)
        btn_layout.addWidget(self.export_excel_btn)

        self.import_json_btn = QPushButton("📥 Импорт JSON")
        self.import_json_btn.clicked.connect(self.import_json)
        btn_layout.addWidget(self.import_json_btn)

        self.export_json_btn = QPushButton("📤 Экспорт JSON")
        self.export_json_btn.clicked.connect(self.export_json)
        btn_layout.addWidget(self.export_json_btn)

        layout.addLayout(btn_layout)

        self.load_data()

    def load_data(self):
        clients = db.get_clients()
        self.table.setRowCount(len(clients))
        for row, c in enumerate(clients):
            self.table.setItem(row, 0, QTableWidgetItem(str(c['id'])))
            self.table.setItem(row, 1, QTableWidgetItem("Физ" if c['type'] == 'individual' else "Юр"))
            self.table.setItem(row, 2, QTableWidgetItem(c['last_name']))
            self.table.setItem(row, 3, QTableWidgetItem(c['first_name']))
            self.table.setItem(row, 4, QTableWidgetItem(c['middle_name']))
            self.table.setItem(row, 5, QTableWidgetItem(c['organization_name']))
            self.table.setItem(row, 6, QTableWidgetItem(c['phone1']))
            self.table.setItem(row, 7, QTableWidgetItem(c['phone2']))
            self.table.setItem(row, 8, QTableWidgetItem(c['email']))
            self.table.setItem(row, 9, QTableWidgetItem(c['messenger']))
        self.table.resizeRowsToContents()

    def get_selected_id(self):
        row = self.table.currentRow()
        if row < 0:
            return None
        id_item = self.table.item(row, 0)
        return int(id_item.text()) if id_item else None

    def add_client(self):
        dialog = ClientEditDialog(self)
        if dialog.exec() == QDialog.DialogCode.Accepted:
            self.load_data()

    def edit_client(self):
        client_id = self.get_selected_id()
        if client_id is None:
            QMessageBox.warning(self, "Ошибка", "Выберите клиента!")
            return
        dialog = ClientEditDialog(self, client_id)
        if dialog.exec() == QDialog.DialogCode.Accepted:
            self.load_data()

    def delete_client(self):
        client_id = self.get_selected_id()
        if client_id is None:
            QMessageBox.warning(self, "Ошибка", "Выберите клиента!")
            return
        reply = QMessageBox.question(self, "Подтверждение", "Удалить клиента и все его адреса?",
                                     QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No)
        if reply == QMessageBox.StandardButton.Yes:
            if db.delete_client(client_id):
                self.load_data()
            else:
                QMessageBox.warning(self, "Ошибка", "Не удалось удалить клиента!")

    def open_addresses(self):
        client_id = self.get_selected_id()
        if client_id is None:
            QMessageBox.warning(self, "Ошибка", "Выберите клиента!")
            return
        dialog = AddressDialog(client_id, self)
        dialog.exec()

    def import_excel(self):
        file_path, _ = QFileDialog.getOpenFileName(self, "Выберите Excel файл", "", "Excel files (*.xlsx *.xls)")
        if file_path:
            try:
                from openpyxl import load_workbook
                wb = load_workbook(file_path)
                ws = wb.active
                imported = 0
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if not row[0] and not row[2]:
                        continue
                    type_str = row[0] or 'individual'
                    type_val = 'individual' if type_str.lower().startswith('физ') else 'legal'
                    last_name = str(row[1]) if row[1] else ''
                    first_name = str(row[2]) if row[2] else ''
                    middle_name = str(row[3]) if row[3] else ''
                    org_name = str(row[4]) if row[4] else ''
                    phone1 = str(row[5]) if row[5] else ''
                    phone2 = str(row[6]) if row[6] else ''
                    email = str(row[7]) if row[7] else ''
                    messenger = str(row[8]) if row[8] else ''
                    db.add_client(type_val, last_name, first_name, middle_name, org_name, phone1, phone2, email, messenger)
                    imported += 1
                QMessageBox.information(self, "Готово", f"Импортировано {imported} клиентов.")
                self.load_data()
            except Exception as e:
                QMessageBox.warning(self, "Ошибка импорта", str(e))

    def export_excel(self):
        file_path, _ = QFileDialog.getSaveFileName(self, "Сохранить Excel", "clients.xlsx", "Excel files (*.xlsx)")
        if file_path:
            try:
                from openpyxl import Workbook
                wb = Workbook()
                ws = wb.active
                ws.title = "Клиенты"
                headers = ["Тип", "Фамилия", "Имя", "Отчество", "Организация",
                           "Телефон 1", "Телефон 2", "Email", "Мессенджер"]
                ws.append(headers)
                clients = db.get_clients()
                for c in clients:
                    ws.append([
                        "Физ" if c['type'] == 'individual' else "Юр",
                        c['last_name'], c['first_name'], c['middle_name'],
                        c['organization_name'], c['phone1'], c['phone2'],
                        c['email'], c['messenger']
                    ])
                wb.save(file_path)
                QMessageBox.information(self, "Готово", f"Экспортировано {len(clients)} клиентов.")
            except Exception as e:
                QMessageBox.warning(self, "Ошибка экспорта", str(e))

    def import_json(self):
        file_path, _ = QFileDialog.getOpenFileName(self, "Выберите JSON файл", "", "JSON files (*.json)")
        if file_path:
            try:
                with open(file_path, 'r', encoding='utf-8') as f:
                    data = json.load(f)
                imported = 0
                for item in data:
                    db.add_client(
                        type=item.get('type', 'individual'),
                        last_name=item.get('last_name', ''),
                        first_name=item.get('first_name', ''),
                        middle_name=item.get('middle_name', ''),
                        organization_name=item.get('organization_name', ''),
                        phone1=item.get('phone1', ''),
                        phone2=item.get('phone2', ''),
                        email=item.get('email', ''),
                        messenger=item.get('messenger', '')
                    )
                    imported += 1
                QMessageBox.information(self, "Готово", f"Импортировано {imported} клиентов.")
                self.load_data()
            except Exception as e:
                QMessageBox.warning(self, "Ошибка импорта", str(e))

    def export_json(self):
        file_path, _ = QFileDialog.getSaveFileName(self, "Сохранить JSON", "clients.json", "JSON files (*.json)")
        if file_path:
            try:
                clients = db.get_clients()
                with open(file_path, 'w', encoding='utf-8') as f:
                    json.dump(clients, f, ensure_ascii=False, indent=2, default=str)
                QMessageBox.information(self, "Готово", f"Экспортировано {len(clients)} клиентов.")
            except Exception as e:
                QMessageBox.warning(self, "Ошибка экспорта", str(e))


class ClientEditDialog(QDialog):
    def __init__(self, parent=None, client_id=None):
        super().__init__(parent)
        self.client_id = client_id
        self.setWindowTitle("Редактирование клиента" if client_id else "Новый клиент")
        self.setMinimumWidth(450)

        layout = QVBoxLayout(self)
        form_layout = QFormLayout()

        self.type_combo = QComboBox()
        self.type_combo.addItems(["Физ. лицо", "Юр. лицо"])
        self.type_combo.currentTextChanged.connect(self.toggle_org_fields)
        form_layout.addRow("Тип клиента:", self.type_combo)

        self.last_name_edit = QLineEdit()
        form_layout.addRow("Фамилия:", self.last_name_edit)

        self.first_name_edit = QLineEdit()
        form_layout.addRow("Имя:", self.first_name_edit)

        self.middle_name_edit = QLineEdit()
        form_layout.addRow("Отчество:", self.middle_name_edit)

        self.org_edit = QLineEdit()
        form_layout.addRow("Организация:", self.org_edit)

        self.phone1_edit = QLineEdit()
        self.phone1_edit.setPlaceholderText("+7 900 123-45-67")
        form_layout.addRow("Телефон 1:", self.phone1_edit)

        self.phone2_edit = QLineEdit()
        self.phone2_edit.setPlaceholderText("+7 900 123-45-67 (доп.)")
        form_layout.addRow("Телефон 2:", self.phone2_edit)

        self.email_edit = QLineEdit()
        self.email_edit.setPlaceholderText("client@example.com")
        form_layout.addRow("Email:", self.email_edit)

        self.messenger_edit = QLineEdit()
        self.messenger_edit.setPlaceholderText("Telegram, WhatsApp, ...")
        form_layout.addRow("Мессенджер:", self.messenger_edit)

        layout.addLayout(form_layout)

        btn_layout = QHBoxLayout()
        self.save_btn = QPushButton("💾 Сохранить")
        self.save_btn.clicked.connect(self.save)
        btn_layout.addWidget(self.save_btn)

        self.cancel_btn = QPushButton("❌ Отмена")
        self.cancel_btn.clicked.connect(self.reject)
        btn_layout.addWidget(self.cancel_btn)

        layout.addLayout(btn_layout)

        if client_id:
            self.load_data()
        self.toggle_org_fields(self.type_combo.currentText())

    def toggle_org_fields(self, text):
        is_legal = text == "Юр. лицо"
        self.org_edit.setEnabled(is_legal)
        self.last_name_edit.setEnabled(not is_legal)
        self.first_name_edit.setEnabled(not is_legal)
        self.middle_name_edit.setEnabled(not is_legal)

    def load_data(self):
        client = db.get_client_by_id(self.client_id)
        if client:
            self.type_combo.setCurrentText("Юр. лицо" if client['type'] == 'legal' else "Физ. лицо")
            self.last_name_edit.setText(client['last_name'])
            self.first_name_edit.setText(client['first_name'])
            self.middle_name_edit.setText(client['middle_name'])
            self.org_edit.setText(client['organization_name'])
            self.phone1_edit.setText(client['phone1'])
            self.phone2_edit.setText(client['phone2'])
            self.email_edit.setText(client['email'])
            self.messenger_edit.setText(client['messenger'])

    def save(self):
        type_val = 'legal' if self.type_combo.currentText() == "Юр. лицо" else 'individual'
        last_name = self.last_name_edit.text().strip()
        first_name = self.first_name_edit.text().strip()
        middle_name = self.middle_name_edit.text().strip()
        org_name = self.org_edit.text().strip()
        phone1 = self.phone1_edit.text().strip()
        phone2 = self.phone2_edit.text().strip()
        email = self.email_edit.text().strip()
        messenger = self.messenger_edit.text().strip()

        if type_val == 'individual' and not last_name and not first_name:
            QMessageBox.warning(self, "Ошибка", "Укажите фамилию и имя!")
            return

        if self.client_id:
            success = db.update_client(self.client_id, type_val, last_name, first_name, middle_name,
                                       org_name, phone1, phone2, email, messenger)
        else:
            success = db.add_client(type_val, last_name, first_name, middle_name,
                                    org_name, phone1, phone2, email, messenger)

        if success:
            self.accept()
        else:
            QMessageBox.warning(self, "Ошибка", "Не удалось сохранить клиента!")